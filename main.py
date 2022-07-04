import json
import time

import openpyxl
import pandas as pd
from openpyxl.chart import LineChart, Reference

from fake import sensors, switchers

BIO_CYCLE = [
    {'period': 2, 'light': 1000, 'temperature': 16, 'moisture': 55, 'humidity': 48},
    {'period': 1, 'light': 500, 'temperature': 18, 'moisture': 60, 'humidity': 50},
    {'period': 1, 'light': 0, 'temperature': 15, 'moisture': 65, 'humidity': 55},
    {'period': 1, 'light': 500, 'temperature': 14, 'moisture': 60, 'humidity': 60}
]

sensors_base = {
    'light': [sensors.light, switchers.light],
    'temperature': [sensors.temperature, switchers.temperature],
    'moisture': [sensors.moisture, switchers.moisture],
    'humidity': [sensors.humidity, switchers.humidity]
}

data_for_cycle = [[cycle.get(name) for name in sensors_base] for cycle in BIO_CYCLE]

time_periods = [period.get("period") for period in BIO_CYCLE]


class DeviceOperator:
    def __init__(self, sensor, switcher):
        self.sensor = sensor
        self.switcher = switcher

    def check_and_control(self, initial_value):

        if initial_value > self.sensor.get_value():
            self.switcher.turn_on()
            return self.sensor.get_value()
        else:
            self.switcher.turn_off()
            return self.sensor.get_value()


class Chart:
    shift_chart = 1

    def __init__(self, column):
        self.column = column

    def plot_chart(self):
        for num in range(self.column):
            chart_type = LineChart()
            col_and_row = Reference(sheet, min_col=num + 1, min_row=1, max_col=num + 1, max_row=sheet.max_row)
            chart_type.add_data(col_and_row, titles_from_data=True)
            sheet.add_chart(chart_type, "L" + str(self.shift_chart))
            self.shift_chart += 15


def writing_reading(log):
    if log == "open":
        with open('E:\Python\Учеба\home works\smart_plant1\data.json', 'r') as data:
            log = json.load(data)
            return log
    elif log == log:
        with open('E:\Python\Учеба\home works\smart_plant1\data.json', 'w') as data:
            json.dump(log, data, indent=2)


def control(number_repetitions):
    for i in range(number_repetitions):
        for number_cycle, data in enumerate(data_for_cycle):
            cycle_change_time = time.time() + time_periods[number_cycle]
            while time.time() < cycle_change_time:
                for number_parameter, (parameter, operator) in enumerate(sensors_base.items()):
                    parameter_control = DeviceOperator(sensors_base.get(parameter)[0], sensors_base.get(parameter)[1])
                    log.setdefault(parameter, [])
                    data = parameter_control.check_and_control(data_for_cycle[number_cycle][number_parameter])
                    log.get(parameter).append(data)
                time.sleep(1)


log = writing_reading("open")

control(int(input("Введите количество повторений био-цикла: ")))

writing_reading(log)

log_xl = pd.DataFrame(log)
log_xl.to_excel("log.xlsx", index=False)
data = openpyxl.load_workbook("log.xlsx")

sheet = data.active

result = Chart(len(sensors_base))
result.plot_chart()
data.save("log.xlsx")


